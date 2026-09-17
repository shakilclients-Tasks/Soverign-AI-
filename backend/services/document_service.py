"""Safe local document ingestion and text extraction."""

import hashlib
import re
import uuid
from pathlib import Path
from typing import Any, Dict, List, Optional

import fitz

from backend.core.config import settings
from backend.core.logging import get_logger
from backend.db.database import get_db_connection
from backend.services.conversation_service import utc_now
from backend.services.ocr_service import ocr_service


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".png",
    ".jpg",
    ".jpeg",
    ".tif",
    ".tiff",
    ".bmp",
    ".webp",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
logger = get_logger("sovereign.documents")


def safe_filename(value: str) -> str:
    name = Path(value).name
    clean = re.sub(r"[^A-Za-z0-9._ -]", "_", name).strip(" .")
    return clean[:180] or "attachment"


class DocumentService:
    def _extract(self, path: Path, extension: str) -> tuple[List[Dict[str, Any]], str]:
        pages: List[Dict[str, Any]] = []
        status = "completed"
        if extension == ".pdf":
            document = fitz.open(path)
            try:
                for index, page in enumerate(document):
                    text = page.get_text("text").strip()
                    image_count = len(page.get_images(full=True))
                    if not text or (image_count > 0 and len(text) < 40):
                        try:
                            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                            ocr_text = ocr_service.extract_image(pixmap.tobytes("png"))
                            if len(ocr_text) > len(text):
                                text = ocr_text
                            if not text:
                                status = "no_text_detected"
                        except RuntimeError as exc:
                            logger.warning("OCR unavailable for PDF page %s: %s", index + 1, exc)
                            status = "ocr_unavailable"
                    pages.append(
                        {
                            "page_number": index + 1,
                            "text": text,
                            "char_count": len(text),
                            "word_count": len(text.split()),
                            "image_count": image_count,
                        }
                    )
            finally:
                document.close()
        elif extension in IMAGE_EXTENSIONS:
            text = ocr_service.extract_image(path.read_bytes())
            if not text:
                status = "no_text_detected"
            pages.append(
                {
                    "page_number": 1,
                    "text": text,
                    "char_count": len(text),
                    "word_count": len(text.split()),
                    "image_count": 1,
                }
            )
        elif extension == ".docx":
            try:
                import docx

                doc = docx.Document(path)
                elements: List[str] = []
                for p in doc.paragraphs:
                    p_text = p.text.strip()
                    if p_text:
                        elements.append(p_text)
                for table in doc.tables:
                    rows_md: List[str] = []
                    for row_idx, row in enumerate(table.rows):
                        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                        rows_md.append("| " + " | ".join(cells) + " |")
                        if row_idx == 0:
                            rows_md.append("| " + " | ".join(["---"] * len(cells)) + " |")
                    if rows_md:
                        elements.append("\n" + "\n".join(rows_md) + "\n")
                text = "\n\n".join(elements).strip()
                if not text:
                    status = "no_text_detected"
                pages.append(
                    {
                        "page_number": 1,
                        "text": text,
                        "char_count": len(text),
                        "word_count": len(text.split()),
                        "image_count": 0,
                    }
                )
            except Exception as exc:
                logger.error("Failed to parse Word docx: %s", exc)
                status = "extraction_failed"
                pages.append(
                    {
                        "page_number": 1,
                        "text": "",
                        "char_count": 0,
                        "word_count": 0,
                        "image_count": 0,
                    }
                )
        elif extension == ".json":
            raw = path.read_bytes()
            try:
                text_raw = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text_raw = raw.decode("latin-1")
            try:
                import json

                parsed = json.loads(text_raw)
                text = json.dumps(parsed, indent=2)
            except Exception:
                text = text_raw.strip()
            pages.append(
                {
                    "page_number": 1,
                    "text": text,
                    "char_count": len(text),
                    "word_count": len(text.split()),
                    "image_count": 0,
                }
            )
        else:
            raw = path.read_bytes()
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("latin-1")
            text = text.strip()
            pages.append(
                {
                    "page_number": 1,
                    "text": text,
                    "char_count": len(text),
                    "word_count": len(text.split()),
                    "image_count": 0,
                }
            )
        return pages, status

    def ingest(self, content: bytes, filename: str, content_type: str) -> Dict[str, Any]:
        clean_name = safe_filename(filename)
        extension = Path(clean_name).suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise ValueError("Supported files are PDF, Word (DOCX), TXT, MD, CSV, JSON, and common images.")
        digest = hashlib.sha256(content).hexdigest()
        with get_db_connection() as conn:
            existing = conn.execute(
                "SELECT * FROM documents WHERE sha256_hash = ?", (digest,)
            ).fetchone()
        if existing:
            result = dict(existing)
            result["duplicate"] = True
            return result

        document_id = str(uuid.uuid4())
        stored_path = (settings.uploads_dir / f"{document_id}_{clean_name}").resolve()
        if not stored_path.is_relative_to(settings.uploads_dir.resolve()):
            raise ValueError("Invalid attachment path")
        stored_path.write_bytes(content)
        try:
            pages, extraction_status = self._extract(stored_path, extension)
        except Exception:
            stored_path.unlink(missing_ok=True)
            raise

        now = utc_now()
        total_chars = sum(page["char_count"] for page in pages)
        total_words = sum(page["word_count"] for page in pages)
        with get_db_connection() as conn:
            conn.execute(
                """
                INSERT INTO documents(
                    id, filename, stored_path, file_type, file_size_bytes, sha256_hash,
                    page_count, char_count, word_count, extraction_status, created_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    document_id,
                    clean_name,
                    str(stored_path),
                    extension.lstrip("."),
                    len(content),
                    digest,
                    len(pages),
                    total_chars,
                    total_words,
                    extraction_status,
                    now,
                ),
            )
            for page in pages:
                conn.execute(
                    """
                    INSERT INTO document_pages(
                        id, document_id, page_number, text_content, char_count,
                        word_count, image_count, created_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        str(uuid.uuid4()),
                        document_id,
                        page["page_number"],
                        page["text"],
                        page["char_count"],
                        page["word_count"],
                        page["image_count"],
                        now,
                    ),
                )
        return {
            "id": document_id,
            "filename": clean_name,
            "stored_path": str(stored_path),
            "file_type": extension.lstrip("."),
            "file_size_bytes": len(content),
            "sha256_hash": digest,
            "page_count": len(pages),
            "char_count": total_chars,
            "word_count": total_words,
            "extraction_status": extraction_status,
            "created_at": now,
            "duplicate": False,
        }

    def get_document(self, document_id: str) -> Optional[Dict[str, Any]]:
        with get_db_connection() as conn:
            row = conn.execute("SELECT * FROM documents WHERE id = ?", (document_id,)).fetchone()
            if not row:
                return None
            pages = conn.execute(
                "SELECT * FROM document_pages WHERE document_id = ? ORDER BY page_number",
                (document_id,),
            ).fetchall()
        result = dict(row)
        result["pages"] = [dict(page) for page in pages]
        has_multiple_pages = len(pages) > 1
        formatted_parts: List[str] = []
        for page in pages:
            body = (page["text_content"] or "").strip()
            if body:
                if has_multiple_pages:
                    formatted_parts.append(f"--- [Page {page['page_number']}] ---\n{body}")
                else:
                    formatted_parts.append(body)
        result["full_text"] = "\n\n".join(formatted_parts)
        return result

    def delete_document(self, document_id: str) -> bool:
        with get_db_connection() as conn:
            row = conn.execute("SELECT stored_path FROM documents WHERE id = ?", (document_id,)).fetchone()
            if not row:
                return False
            conn.execute("DELETE FROM documents WHERE id = ?", (document_id,))
        path = Path(row["stored_path"]).resolve()
        if path.is_relative_to(settings.uploads_dir.resolve()):
            path.unlink(missing_ok=True)
        return True


document_service = DocumentService()
